# ==========================================
# STUDENT RESULT MANAGEMENT SYSTEM
# ==========================================

from openpyxl import Workbook, load_workbook
import os


# ==========================================
# EXCEL FILE
# ==========================================

EXCEL_FILE = "student_results.xlsx"


# ==========================================
# FUNCTION: CREATE EXCEL FILE
# ==========================================

def create_excel_file():

    if not os.path.exists(EXCEL_FILE):

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Student Results"

        # Excel column headings
        sheet.append([
            "Roll No",
            "Name",
            "Class",
            "Subject 1",
            "Subject 2",
            "Subject 3",
            "Subject 4",
            "Subject 5",
            "Total",
            "Percentage",
            "Grade",
            "Status"
        ])

        workbook.save(EXCEL_FILE)


# ==========================================
# FUNCTION: CALCULATE RESULT
# ==========================================

def calculate_result(marks, total_marks):

    total = sum(marks)

    percentage = (total / (total_marks * 5)) * 100

    # Grade
    if percentage >= 90:
        grade = "A+"

    elif percentage >= 80:
        grade = "A"

    elif percentage >= 70:
        grade = "B"

    elif percentage >= 60:
        grade = "C"

    elif percentage >= 50:
        grade = "D"

    else:
        grade = "F"

    # Pass / Fail
    if percentage >= 40:
        status = "PASS"

    else:
        status = "FAIL"

    return total, percentage, grade, status


# ==========================================
# FUNCTION 1: ADD STUDENT
# ==========================================

def add_student():

    print("\n")
    print("==========================================")
    print("          ADD STUDENT RESULT")
    print("==========================================")

    # --------------------------------------
    # Roll Number
    # --------------------------------------

    while True:

        try:
            roll_no = int(input("Enter student roll no: "))

            if roll_no > 0:
                break

            else:
                print("Roll number must be greater than 0.")

        except ValueError:
            print("Please enter a valid roll number.")


    # --------------------------------------
    # Check duplicate roll number
    # --------------------------------------

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if row[0] == roll_no:

            print("\nStudent with this roll number already exists.")
            workbook.close()
            return

    workbook.close()


    # --------------------------------------
    # Student Name
    # --------------------------------------

    while True:

        name = input("Enter student name: ").strip()

        if name != "":
            break

        print("Name cannot be empty.")


    # --------------------------------------
    # Student Class / Course
    # --------------------------------------

    while True:

        stud_class = input("Enter student class/course: ").strip()

        if stud_class != "":
            break

        print("Class/Course cannot be empty.")


    # --------------------------------------
    # Total Marks Per Subject
    # --------------------------------------

    while True:

        try:

            total_marks = int(
                input("Enter total marks per subject: ")
            )

            if total_marks > 0:
                break

            else:
                print("Total marks must be greater than 0.")

        except ValueError:
            print("Please enter a valid number.")


    # ======================================
    # ENTER MARKS OF 5 SUBJECTS
    # ======================================

    marks = []

    print("\nEnter marks of 5 subjects:")

    for i in range(1, 6):

        while True:

            try:

                mark = int(
                    input(f"Enter marks for Subject {i}: ")
                )

                if 0 <= mark <= total_marks:

                    marks.append(mark)
                    break

                else:

                    print(
                        "Invalid marks! Enter marks between 0 and",
                        total_marks
                    )

            except ValueError:

                print("Please enter a valid number.")


    # ======================================
    # CALCULATE RESULT
    # ======================================

    total, percentage, grade, status = calculate_result(
        marks,
        total_marks
    )


    # ======================================
    # SAVE DATA TO EXCEL
    # ======================================

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    sheet.append([
        roll_no,
        name,
        stud_class,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        round(percentage, 2),
        grade,
        status
    ])

    workbook.save(EXCEL_FILE)
    workbook.close()


    # ======================================
    # SUCCESS MESSAGE
    # ======================================

    print("\n==========================================")
    print("     STUDENT RESULT ADDED SUCCESSFULLY")
    print("==========================================")

    print("Total      :", total)
    print("Percentage :", round(percentage, 2), "%")
    print("Grade      :", grade)
    print("Status     :", status)


# ==========================================
# FUNCTION 2: GET STUDENT RESULT
# ==========================================

def get_result():

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    # --------------------------------------
    # Check if students exist
    # --------------------------------------

    if sheet.max_row <= 1:

        print("\nNO STUDENT DATA FOUND.")
        workbook.close()
        return


    # --------------------------------------
    # Enter Roll Number
    # --------------------------------------

    while True:

        try:

            roll_no = int(
                input("\nEnter student roll no: ")
            )

            break

        except ValueError:

            print("Please enter a valid roll number.")


    # --------------------------------------
    # Search student in Excel
    # --------------------------------------

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        if row[0] == roll_no:

            print("\n")
            print("==========================================")
            print("              STUDENT RESULT")
            print("==========================================")

            print("Roll No     :", row[0])
            print("Name        :", row[1])
            print("Class       :", row[2])

            print("------------------------------------------")

            print("Subject 1   :", row[3])
            print("Subject 2   :", row[4])
            print("Subject 3   :", row[5])
            print("Subject 4   :", row[6])
            print("Subject 5   :", row[7])

            print("------------------------------------------")

            print("Total       :", row[8])
            print("Percentage  :", row[9], "%")
            print("Grade       :", row[10])
            print("Status      :", row[11])

            print("==========================================")

            workbook.close()
            return


    # --------------------------------------
    # Student Not Found
    # --------------------------------------

    print("\nNo student found with roll no:", roll_no)

    workbook.close()


# ==========================================
# FUNCTION 3: SHOW ALL DATA
# ==========================================

def show_all_data():

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    # --------------------------------------
    # Check if data exists
    # --------------------------------------

    if sheet.max_row <= 1:

        print("\nNO STUDENT DATA FOUND.")
        workbook.close()
        return


    print("\n")
    print("=" * 100)
    print("                     ALL STUDENT DATA")
    print("=" * 100)


    # --------------------------------------
    # Table Heading
    # --------------------------------------

    print(
        f"{'Roll No':<10}"
        f"{'Name':<15}"
        f"{'Class':<12}"
        f"{'Total':<10}"
        f"{'Percentage':<14}"
        f"{'Grade':<10}"
        f"{'Status':<10}"
    )

    print("-" * 100)


    # --------------------------------------
    # Read all records from Excel
    # --------------------------------------

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        print(
            f"{str(row[0]):<10}"
            f"{str(row[1]):<15}"
            f"{str(row[2]):<12}"
            f"{str(row[8]):<10}"
            f"{str(row[9]):<14}"
            f"{str(row[10]):<10}"
            f"{str(row[11]):<10}"
        )


    print("=" * 100)

    workbook.close()


# ==========================================
# FUNCTION 4: MENU
# ==========================================

def menu():

    while True:

        print("\n")
        print("=" * 50)
        print("       STUDENT RESULT MANAGEMENT SYSTEM")
        print("=" * 50)

        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")

        print("=" * 50)


        # --------------------------------------
        # Take Choice
        # --------------------------------------

        choice = input("Enter your choice: ")


        # --------------------------------------
        # OPTION 1
        # --------------------------------------

        if choice == "1":

            add_student()


        # --------------------------------------
        # OPTION 2
        # --------------------------------------

        elif choice == "2":

            get_result()


        # --------------------------------------
        # OPTION 3
        # --------------------------------------

        elif choice == "3":

            show_all_data()


        # --------------------------------------
        # OPTION 4
        # --------------------------------------

        elif choice == "4":

            print("\nThank you!")

            break


        # --------------------------------------
        # INVALID CHOICE
        # --------------------------------------

        else:

            print("\nEnter valid choice")


# ==========================================
# START PROGRAM
# ==========================================

create_excel_file()

menu()
